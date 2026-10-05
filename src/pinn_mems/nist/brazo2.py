"""Brazo 2: trazas de deformación residual (vigas biempotradas) y de gradiente
de deformación (voladizos) de NIST SP 260-177.

Funciones extraídas del notebook `notebooks/EDA_Brazo2_NIST.ipynb`.
"""

from __future__ import annotations

import re
from pathlib import Path

import pandas as pd

from pinn_mems.nist.archivos import CARPETA_CRUDOS

DATA_DIR = CARPETA_CRUDOS

# Columnas que identifican una traza
TRACE_KEYS = ["test_type", "material", "trace"]

COLUMN_ORDER = [
    "test_type", "structure", "material", "chip", "length_um",
    "trace", "x", "z", "calibrated", "file",
]

# Ej.: ...Trace.ap.RM.8097.0108.P2.0deg.L500.Ins3.y.239.421.xlsx
_PATRON_NOMBRE = re.compile(
    r"Trace\.(?P<trace>[^.]+)\.RM\.(?P<rm>\d{4})\.(?P<chip>\d{4})\."
    r"(?:(?P<layer>P\d)\.(?P<orientation>\d+)deg\.)?"
    r"L(?P<length>\d+)\.Ins\d(?:\.y\.(?P<y>[\d.]+))?\.xlsx$"
)


def get_trace_metadata(file_path) -> dict:
    """Extract metadata from a NIST trace filename."""
    filename = Path(file_path).name
    match = _PATRON_NOMBRE.search(filename)
    if match is None:
        raise ValueError(f"Nombre de traza no reconocido: {filename}")

    residual = "RESIDUAL.STRAIN" in filename.upper()
    return {
        "test_type": "Residual Strain" if residual else "Strain Gradient",
        # Deformación residual: vigas biempotradas; gradiente: voladizos (datos/nist/LEEME.md)
        "structure": "fixed-fixed" if residual else "cantilever",
        "material": f"RM {match['rm']}",
        "chip": match["chip"],
        "layer": match["layer"],
        "orientation_deg": int(match["orientation"]) if match["orientation"] else None,
        "length_um": int(match["length"]),
        "y_um": float(match["y"]) if match["y"] else None,
        "trace": match["trace"],
        "file": filename,
    }


def _calibration_factors(df: pd.DataFrame) -> dict:
    """calx y calz de la hoja: el valor va en la primera fila bajo su encabezado."""
    factors = {}
    for column in df.columns:
        if isinstance(column, str) and column.strip().lower().startswith(("calx", "calz")):
            factors[column.strip()[:4].lower()] = float(df[column].iloc[0])
    return factors


def load_trace(file_path, calibrate: bool = True) -> pd.DataFrame:
    """Load x and z measurements (µm) from a NIST Excel trace.

    Uses the x column in micrometers (some sheets also have one in mm). When the
    sheet provides calibration factors (calx, calz) and ``calibrate`` is True,
    returns x·calx and z·calz.
    """
    df = pd.read_excel(file_path)

    def columns_starting_with(*prefixes):
        return [c for c in df.columns if isinstance(c, str) and c.strip().startswith(prefixes)]

    x_columns = columns_starting_with("x-uncal", "x (uncal)")
    z_columns = columns_starting_with("z-uncal", "z (uncal)")
    if not x_columns or not z_columns:
        raise ValueError(f"No se encontraron columnas x/z válidas en {Path(file_path).name}")

    # La segunda fila de la hoja indica las unidades de cada columna
    x_column = next((c for c in x_columns if str(df[c].iloc[1]).strip().lower() == "um"), x_columns[0])
    trace_df = df[[x_column, z_columns[0]]].copy()
    trace_df.columns = ["x", "z"]
    trace_df["x"] = pd.to_numeric(trace_df["x"], errors="coerce")
    trace_df["z"] = pd.to_numeric(trace_df["z"], errors="coerce")
    trace_df = trace_df.dropna(subset=["x", "z"]).reset_index(drop=True)

    factors = _calibration_factors(df) if calibrate else {}
    trace_df["x"] *= factors.get("calx", 1.0)
    trace_df["z"] *= factors.get("calz", 1.0)
    trace_df["calibrated"] = bool(factors)

    for column, value in get_trace_metadata(file_path).items():
        trace_df[column] = value
    return trace_df


def load_all_traces(data_dir: Path = DATA_DIR, calibrate: bool = True) -> pd.DataFrame:
    """Dataset maestro: todas las trazas de deformación de `data_dir`."""
    files = sorted(Path(data_dir).glob("*STRAIN*.xlsx"))
    df = pd.concat([load_trace(f, calibrate=calibrate) for f in files], ignore_index=True)
    return df[COLUMN_ORDER]


# --- Atípicos -----------------------------------------------------------------


def _iqr_bounds(values: pd.Series) -> tuple[float, float, float, float]:
    q1, q3 = values.quantile(0.25), values.quantile(0.75)
    iqr = q3 - q1
    return q1, q3, q1 - 1.5 * iqr, q3 + 1.5 * iqr


def detect_iqr_outliers(group: pd.DataFrame, column: str = "z") -> pd.Series:
    """Count potential outliers using the IQR rule within a trace."""
    q1, q3, lower_bound, upper_bound = _iqr_bounds(group[column])
    outliers = ((group[column] < lower_bound) | (group[column] > upper_bound)).sum()
    return pd.Series({
        "q1": q1,
        "q3": q3,
        "iqr": q3 - q1,
        "lower_bound": lower_bound,
        "upper_bound": upper_bound,
        "outliers": int(outliers),
        "outlier_percentage": 100 * outliers / len(group),
    })


def iqr_outlier_mask(df: pd.DataFrame, column: str = "z", by=TRACE_KEYS) -> pd.Series:
    """Marca (True) los puntos fuera de [Q1 − 1.5·IQR, Q3 + 1.5·IQR] de su traza.

    Solo detecta: no elimina filas.
    """
    def mask(group):
        _, _, lower, upper = _iqr_bounds(group)
        return (group < lower) | (group > upper)

    return df.groupby(by)[column].transform(mask).astype(bool)


# --- Valores de la hoja del NIST y validación -----------------------------------


def read_sheet_parameters(file_path) -> dict:
    """Parámetros que la hoja del NIST trae en su encabezado.

    Siempre que existan: calx, calz, f (offset x1ave·calx del eje v, en µm) y
    alpha (ángulo de desalineación, rad). En gradiente de deformación además
    Rint, m, n (círculo del modelo, µm) y s (signo de la curvatura).
    """
    import openpyxl

    hoja = openpyxl.load_workbook(file_path, data_only=True).active
    etiquetas = {
        "calx": ("calx",), "calz": ("calz",), "f": ("f=", "f ="), "alpha": ("α",),
        "Rint": ("rint",), "m": ("m",), "n": ("n",), "s": ("s",),
    }
    params = {}
    for fila in hoja.iter_rows(min_row=1, max_row=12):
        for celda in fila:
            if not isinstance(celda.value, str):
                continue
            texto = celda.value.strip().lower()
            for clave, prefijos in etiquetas.items():
                if clave in params:
                    continue
                exacta = clave in ("m", "n", "s")
                if (texto in prefijos) if exacta else texto.startswith(prefijos):
                    # calx y calz llevan el valor debajo; el resto, a la derecha
                    fila_v, col_v = (celda.row + 1, celda.column) if clave in ("calx", "calz") else (celda.row, celda.column + 1)
                    valor = hoja.cell(row=fila_v, column=col_v).value
                    if isinstance(valor, (int, float)):
                        params[clave] = float(valor)
    # m, n y s solo tienen este significado junto al círculo (hojas de gradiente)
    if "Rint" not in params:
        for clave in ("m", "n", "s"):
            params.pop(clave, None)
    return params


def v_axis(x_uncal, calx: float, alpha: float, f: float):
    """Posición a lo largo de la viga (Ecs. RS10–RS14 / SG del SP 260-177):
    v = (x·calx − f)·cos α + f."""
    import numpy as np

    return (np.asarray(x_uncal, dtype=float) * calx - f) * np.cos(alpha) + f


def read_nist_columns(file_path) -> pd.DataFrame:
    """Columnas que calcula la propia hoja del NIST: v, z calibrada y modelo(s)."""
    import openpyxl

    hoja = openpyxl.load_workbook(file_path, data_only=True).active
    # Puede haber más de un rótulo "v-axis data": vale el que tiene z a su derecha
    inicio = next(
        (c for fila in hoja.iter_rows(max_row=12) for c in fila
         if isinstance(c.value, str) and c.value.strip().startswith("v-axis")
         and str(hoja.cell(row=c.row, column=c.column + 1).value).strip().lower().startswith("z")),
        None,
    )
    if inicio is None:
        raise ValueError(f"La hoja no trae columnas calculadas por el NIST: {Path(file_path).name}")
    nombres = []
    for k in range(4):
        valor = hoja.cell(row=inicio.row, column=inicio.column + k).value
        if not isinstance(valor, str):
            break
        nombres.append(valor.strip())
    filas = []
    for r in range(inicio.row + 1, hoja.max_row + 1):
        valores = [hoja.cell(row=r, column=inicio.column + k).value for k in range(len(nombres))]
        if isinstance(valores[0], (int, float)) and isinstance(valores[1], (int, float)):
            filas.append([v if isinstance(v, (int, float)) else float("nan") for v in valores])
    return pd.DataFrame(filas, columns=nombres)


def circle_through_points(points) -> tuple[float, float, float]:
    """Círculo que pasa por tres puntos (v, z): devuelve (Rint, m, n)."""
    import numpy as np

    (x1, y1), (x2, y2), (x3, y3) = points
    A = 2 * np.array([[x2 - x1, y2 - y1], [x3 - x1, y3 - y1]], dtype=float)
    b = np.array([x2**2 - x1**2 + y2**2 - y1**2, x3**2 - x1**2 + y3**2 - y1**2])
    m, n = np.linalg.solve(A, b)
    return float(np.hypot(x1 - m, y1 - n)), float(m), float(n)


def strain_gradient(Rint_um: float, s: float) -> float:
    """Gradiente de deformación en m⁻¹ a partir del radio Rint (µm).

    Convención del SP 260-177: s = −1 para voladizos que se curvan hacia
    arriba, con gradiente positivo (ejemplo resuelto, p. 190).
    """
    return -s / (Rint_um * 1e-6)


# --- Exportación a CSV -----------------------------------------------------------

_ESTRUCTURA_ES = {"fixed-fixed": "biempotrada", "cantilever": "voladizo"}


def trace_csv_name(file_path) -> str:
    """p. ej. biempotrada_RM8096-0009_L200_traza_b.csv"""
    meta = get_trace_metadata(file_path)
    material = meta["material"].replace(" ", "")
    return f"{_ESTRUCTURA_ES[meta['structure']]}_{material}-{meta['chip']}_L{meta['length_um']}_traza_{meta['trace']}.csv"


def trace_for_export(file_path) -> pd.DataFrame:
    """Tabla de una traza con las columnas de datos/nist/trazas/."""
    import numpy as np

    trace = load_trace(file_path)
    crudo = load_trace(file_path, calibrate=False)
    params = read_sheet_parameters(file_path)
    if {"calx", "alpha", "f"} <= params.keys():
        v = v_axis(crudo["x"], params["calx"], params["alpha"], params["f"])
    else:
        v = np.full(len(trace), np.nan)
    meta = get_trace_metadata(file_path)
    return pd.DataFrame({
        "x_um": trace["x"],
        "z_um": trace["z"],
        "v_um": v,
        "traza": meta["trace"],
        "estructura": _ESTRUCTURA_ES[meta["structure"]],
        "material": meta["material"],
        "chip": meta["chip"],
        "L_um": meta["length_um"],
        "calibrada": trace["calibrated"],
        "archivo_origen": meta["file"],
    })


# --- Perfil sobre la viga contra el modelo del NIST (T52) ----------------------


def beam_profile(file_path, edge_fraction: float = 0.2, jump_um: float = 0.3) -> pd.DataFrame:
    """Tramo de la viga que modela el NIST: v, z calibrada, modelo y residuo (µm).

    Solo existe en las trazas a lo largo de la viga (b, c, d). En deformación
    residual el modelo son dos cosenos (zmodel1 y zmodel2, unidos en el pico);
    en gradiente de deformación, un círculo.

    - `xi` va de 0 a 1 a lo largo del tramo, medido desde el borde del anclaje,
      que la hoja da como f = x1ave·calx (en orientación de 180° está en |v| grande).
    - `on_beam` es False en las mesetas de los extremos que quedan sobre el soporte:
      se detectan por un cambio de nivel sostenido de z mayor que `jump_um` en la
      fracción `edge_fraction` de cada extremo. El modelo no aplica ahí.
    """
    import numpy as np

    nist = read_nist_columns(file_path)
    v = nist.iloc[:, 0].to_numpy()
    z = nist.iloc[:, 1].to_numpy()
    modelo = nist.iloc[:, 2:].bfill(axis=1).iloc[:, 0].to_numpy()
    f = read_sheet_parameters(file_path)["f"]
    distancia = np.abs(v - f)
    xi = (distancia - distancia.min()) / np.ptp(distancia)

    # Salto sostenido: cambio de nivel entre las medianas de `ventana` puntos a cada
    # lado (un pico aislado de ruido no cuenta)
    on_beam = np.ones(len(z), dtype=bool)
    ventana = 6
    saltos = [
        k for k in range(ventana, len(z) - ventana)
        if abs(np.median(z[k + 1:k + 1 + ventana]) - np.median(z[k + 1 - ventana:k + 1])) > jump_um
        and abs(z[k + 1] - z[k]) > jump_um
    ]
    for k in saltos:
        if k < edge_fraction * len(z):
            on_beam[: k + 1] = False
        elif k >= (1 - edge_fraction) * len(z):
            on_beam[k + 1:] = False

    meta = get_trace_metadata(file_path)
    return pd.DataFrame({
        "v_um": v, "xi": xi, "z_um": z, "z_model_um": modelo, "residual_um": z - modelo,
        "on_beam": on_beam,
        "structure": meta["structure"], "material": meta["material"], "length_um": meta["length_um"],
        "trace": meta["trace"], "file": meta["file"],
    })


def residual_by_zone(profile: pd.DataFrame, edge: float = 0.15) -> dict:
    """RMS del residuo cerca de cada extremo del tramo y en el centro, y su autocorrelación.

    Usa solo los puntos sobre la viga (`on_beam`) cuando la columna existe.

    Una autocorrelación de un paso cercana a 1 indica residuos con estructura
    (desajuste del modelo o ruido correlacionado), no ruido blanco.
    """
    import numpy as np

    if "on_beam" in profile:
        profile = profile[profile["on_beam"]]
    r = profile["residual_um"].to_numpy()
    xi = profile["xi"].to_numpy()
    rms = lambda m: float(np.sqrt(np.mean(r[m] ** 2)))  # noqa: E731
    return {
        "rms_start_um": rms(xi <= edge),
        "rms_center_um": rms((xi > edge) & (xi < 1 - edge)),
        "rms_end_um": rms(xi >= 1 - edge),
        "lag1_autocorrelation": float(np.corrcoef(r[:-1], r[1:])[0, 1]),
    }
